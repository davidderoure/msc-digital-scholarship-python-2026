"""Build slides/python_syntax_slides.pptx (MSc Digital Scholarship 2026).

House style for classroom projection: white backgrounds throughout, large type
(body 24pt, code 20-24pt, titles 40pt, footer 14pt), Cambria headings,
Calibri body, Courier New code.

Usage:  python tools/build_syntax_slides.py [output.pptx]
Needs:  pip install python-pptx
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

COURSE = "MSc Digital Scholarship 2026"
DARK = RGBColor(0x1F, 0x3A, 0x34)     # deep green
ACCENT = RGBColor(0xB5, 0x45, 0x2B)   # rubric red
BLUE = RGBColor(0x2E, 0x6F, 0x8E)
GOLD = RGBColor(0x8A, 0x67, 0x0E)
TINT = RGBColor(0xEE, 0xF3, 0xF0)     # code panels and cards
MUTED = RGBColor(0x55, 0x60, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"

# Type scale
TITLE, HUGE, SUB = 40, 54, 28        # slide title, title-slide title, subheading
TEXT, SMALL = 24, 20                 # body text, secondary text
CODE, CODE_S = 24, 20                # code in narrow and wide panels
FOOT = 14

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def textbox(slide, x, y, w, h, paras, size=TEXT, font=BODY, color=INK, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, margin=0.05, space_after=6):
    """paras: list of paragraphs; each paragraph is a str or a list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for text, o in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", False)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def panel(slide, x, y, w, h, fill=TINT, radius=0.06):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def code_runs(line):
    """Grey comments, red '→ output' lines."""
    if line.lstrip().startswith("→"):
        return [(line, {"color": ACCENT})]
    if "#" in line:
        i = line.index("#")
        return [(line[:i], {}), (line[i:], {"color": MUTED, "italic": True})]
    return [(line or " ", {})]


def code(slide, x, y, w, h, src, size=CODE):
    panel(slide, x, y, w, h)
    lines = src.strip("\n").split("\n")
    textbox(slide, x + 0.2, y + 0.15, w - 0.4, h - 0.3, [code_runs(l) for l in lines],
            size=size, font=MONO, space_after=2)


def footer(s):
    textbox(s, 0.6, 6.95, 6, 0.4, [COURSE], size=FOOT, color=MUTED, margin=0)
    textbox(s, W - 1.6, 6.95, 1.0, 0.4, [str(len(prs.slides))], size=FOOT, color=MUTED,
            align=PP_ALIGN.RIGHT, margin=0)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def new_slide(title, notes_text=""):
    s = prs.slides.add_slide(BLANK)
    textbox(s, 0.6, 0.35, W - 1.2, 0.9, [title], size=TITLE, font=HEAD, color=DARK,
            anchor=MSO_ANCHOR.MIDDLE, margin=0)
    footer(s)
    if notes_text:
        notes(s, notes_text)
    return s


# 1 ─ Title ───────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
textbox(s, 0.9, 1.4, W - 1.8, 2.6, ["Python: the grammar behind", "what you've been doing"], size=HUGE,
        font=HEAD, color=DARK, anchor=MSO_ANCHOR.BOTTOM, margin=0, space_after=0)
textbox(s, 0.9, 4.3, W - 1.8, 1.0, ["Syntax and constructs, from values to functions"], size=SUB, margin=0)
footer(s)
notes(s, "So far we've learned Python by using it: immersion. Now the grammar book. "
         "Nothing on these slides is new; it's the system behind what you've already typed.")

# 2 ─ Why now ─────────────────────────────────────────────────────────────
s = new_slide("Two ways to learn a language",
              "Analogy with learning a natural language. Most of you learned a second language with both: "
              "immersion and a grammar book. Each makes sense of the other. "
              "The notebooks were immersion. Today is the grammar.")
for i, (head, body, eg) in enumerate([
    ("Immersion", "What we did in Colab: run it, change it, see what happens.",
     "words = sample.split()\nwords[-3:]\n→ ['knew', 'by', 'noon.']"),
    ("The grammar book", "Today: name the parts, see the patterns.",
     "thing.method(argument)\nlist[start:stop]"),
]):
    x = 0.6 + i * 6.25
    textbox(s, x, 1.5, 5.9, 0.6, [head], size=32, font=HEAD, color=ACCENT, margin=0)
    textbox(s, x, 2.2, 5.9, 1.2, [body], margin=0)
    code(s, x, 3.5, 5.9, 2.1, eg)
textbox(s, 0.6, 5.9, 12, 0.7, [[("Both help. Each makes sense of the other.", {"italic": True, "color": DARK})]],
        size=SUB, margin=0)

# 3 ─ Values and types ───────────────────────────────────────────────────
s = new_slide("Every value has a type",
              "Python cares about types because each type has its own operations. "
              "Try type() on anything you're unsure about. Common confusion: '10' (a string) vs 10 (a number).")
rows = [("str", "text", '"Bennet"'), ("int", "whole number", "39"),
        ("float", "decimal number", "4.2"), ("list", "ordered collection", "['Mr.', 'Bennet']"),
        ("bool", "true or false", "True")]
for i, (t, desc, eg) in enumerate(rows):
    y = 1.5 + i * 1.0
    panel(s, 0.6, y, 1.5, 0.8, fill=DARK)
    textbox(s, 0.6, y, 1.5, 0.8, [t], size=CODE, font=MONO, color=WHITE, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 2.3, y, 2.6, 0.8, [desc], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 4.9, y, 3.6, 0.8, [eg], size=CODE, font=MONO, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
code(s, 8.7, 1.5, 4.03, 4.8, '''
type("Bennet")
→ str

"10" + "5"
→ '105'

10 + 5
→ 15
''')

# 4 ─ Variables ──────────────────────────────────────────────────────────
s = new_slide("= gives a name; == asks a question",
              "The single most common beginner confusion. '=' is assignment: make this name refer to this value. "
              "'==' is comparison: are these equal? Answer True or False. Capitals matter: 'Bennet' is not 'bennet'.")
textbox(s, 0.6, 1.5, 5.8, 4.8, [
    [("=", {"font": MONO, "bold": True, "color": ACCENT, "size": 36}), ("   assignment", {"size": SUB, "bold": True})],
    "Make this name refer to this value.",
    " ",
    [("==", {"font": MONO, "bold": True, "color": ACCENT, "size": 36}), ("  comparison", {"size": SUB, "bold": True})],
    "Are these equal? The answer is True or False.",
], margin=0)
code(s, 6.8, 1.5, 5.93, 4.8, '''
size = len(sample)

size == 196
→ True

name = "Bennet"
name == "bennet"
→ False
''')

# 5 ─ Three ways to call ─────────────────────────────────────────────────
s = new_slide("Three ways to call a tool",
              "The key slide. All three patterns appeared in the notebooks. The difference is what comes before the dot. "
              "re.split(pattern, sample) and sample.split() do similar jobs but are different tools, "
              "so the text goes in a different place.")
cards = [("FUNCTION", "function(thing)", "len(sample)", "do len to sample"),
         ("METHOD", "thing.method()", "sample.split()", "take sample and split it"),
         ("MODULE FUNCTION", "module.function()", "statistics.mean(x)", "mean, from the statistics module")]
for i, (name, pattern, eg, read) in enumerate(cards):
    y = 1.5 + i * 1.45
    panel(s, 0.6, y, 12.13, 1.3)
    textbox(s, 0.85, y + 0.1, 4.5, 0.4, [[(name, {"bold": True, "color": ACCENT})]], size=16, margin=0)
    textbox(s, 0.85, y + 0.5, 4.6, 0.65, [pattern], size=CODE, font=MONO, color=DARK, margin=0)
    textbox(s, 5.6, y + 0.5, 3.9, 0.65, [eg], size=CODE, font=MONO, margin=0)
    textbox(s, 9.6, y + 0.2, 3.0, 1.0, [[(read, {"italic": True})]], size=SMALL, margin=0,
            anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 0.6, 5.95, 12.1, 0.8, [[("sample.split()", {"font": MONO, "color": ACCENT}), ("  vs  ", {}),
                                   ("re.split(pattern, sample)", {"font": MONO, "color": ACCENT}),
                                   (": same job, different tool", {})]], margin=0)

# 6 ─ Arguments and return values ─────────────────────────────────────────
s = new_slide("What goes in, what comes back",
              "Arguments go in the brackets, separated by commas. Most tools give back (return) a value you can store, "
              "e.g. quiet = sample.lower(). print is the odd one out: it shows something but gives back nothing (None). "
              "String methods never change the original string; they return a new one.")
code(s, 0.6, 1.5, 7.0, 5.2, '''
# one argument
len(sample)
→ 196

# two arguments
sample.replace("Bennet", "Darcy")
→ 'Mr. Darcy looked up ...'

# the original is unchanged
sample
→ 'Mr. Bennet looked up ...'
''', size=CODE_S)
textbox(s, 8.0, 1.5, 4.73, 5.2, [
    [("Arguments", {"bold": True, "color": DARK, "size": SUB})],
    "go inside the brackets, separated by commas.",
    " ",
    [("Return values", {"bold": True, "color": DARK, "size": SUB})],
    "are what comes back. Keep one by naming it.",
    " ",
    [("print(x)", {"font": MONO, "color": ACCENT}), (" shows x, but gives back nothing.", {})],
], margin=0)

# 7 ─ Strings ──────────────────────────────────────────────────────────────
s = new_slide("Strings: text and its methods",
              "Single or double quotes, but they must match. Text containing double quotes goes inside single quotes, "
              "and the other way round. These are the string methods we've used; there are many more, and an AI can "
              "tell you about them, but test what it says. startswith and endswith give True or False.")
textbox(s, 0.6, 1.5, 5.4, 0.5, [[("Quotes must match", {"bold": True, "color": DARK})]], size=SUB, margin=0)
code(s, 0.6, 2.15, 5.4, 2.0, '''
'Mr. Bennet'
"my dear?"
'"Is it?" he said.'
''')
textbox(s, 0.6, 4.4, 5.4, 2.0, ["Text with double quotes goes inside single quotes."], margin=0)
methods = [("lower()", "lower case"), ("upper()", "capitals"), ('count("x")', "how many x"),
           ('replace("a", "b")', "swap a for b"), ("split()", "list of words"), ('strip(".,")', "trim the ends")]
for i, (m, d) in enumerate(methods):
    y = 1.5 + i * 0.85
    textbox(s, 6.5, y, 4.0, 0.7, [[("." + m, {"font": MONO, "color": ACCENT})]], size=CODE, margin=0,
            anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 10.6, y, 2.2, 0.7, [d], margin=0, anchor=MSO_ANCHOR.MIDDLE)

# 8 ─ Lists, indexing and slicing ────────────────────────────────────────
s = new_slide("Lists: counting from 0, and from the end",
              "Positions start at 0. Negative numbers count back from the end. "
              "A slice [start:stop] includes start but stops before stop. Leaving a number out means "
              "'from the beginning' or 'to the end'. len(ending) is 6, but the last position is 5. "
              "Exercise from the warm-up: three ways to get the last three words. Which works for any list?")
ending = ["The", "whole", "street", "knew", "by", "noon."]
bw, x0, y0 = 1.85, 0.75, 2.05
for i, w in enumerate(ending):
    x = x0 + i * (bw + 0.12)
    textbox(s, x, y0 - 0.6, bw, 0.5, [str(i)], size=CODE, font=MONO, color=BLUE, align=PP_ALIGN.CENTER, margin=0)
    panel(s, x, y0, bw, 0.9, fill=DARK, radius=0.1)
    textbox(s, x, y0, bw, 0.9, [repr(w)], size=SMALL, font=MONO, color=WHITE, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, margin=0)
    textbox(s, x, y0 + 1.0, bw, 0.5, [str(i - 6)], size=CODE, font=MONO, color=ACCENT, align=PP_ALIGN.CENTER,
            margin=0)
code(s, 0.6, 3.75, 7.6, 2.95, '''
ending[0]     → 'The'
ending[-1]    → 'noon.'
ending[1:3]   → ['whole', 'street']
ending[-3:]   → ['knew', 'by', 'noon.']
''', size=CODE_S)
textbox(s, 8.6, 3.75, 4.13, 2.95, [
    [("[start:stop]", {"font": MONO, "color": ACCENT})],
    [("stops ", {}), ("before", {"italic": True}), (" stop", {})],
    " ",
    [("Leave a number out", {"bold": True})],
    "start or end of the list",
], margin=0, space_after=2)

# 9 ─ List comprehensions ────────────────────────────────────────────────
s = new_slide("List comprehensions",
              "Make a new list from an old one. Read from the middle outwards. The 'if' part is optional. "
              "If you know R: this is the job sapply / Filter, or plain vectorised code, do there. "
              "If you know Lisp: map and filter.")
parts = [("[", INK), ("len(w)", ACCENT), (" ", INK), ("for w in words", BLUE), (" ", INK), ("if len(w) > 5", GOLD), ("]", INK)]
textbox(s, 0.6, 1.4, 12.1, 0.9, [[(t, {"color": c, "bold": True}) for t, c in parts]], size=36, font=MONO,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
labels = [("what to make", "the length of each word", ACCENT),
          ("from where", "each item in turn, called w", BLUE),
          ("which ones", "keep only items that pass (optional)", GOLD)]
for i, (head, body, c) in enumerate(labels):
    x = 0.6 + i * 4.1
    panel(s, x, 2.55, 3.85, 1.75)
    textbox(s, x + 0.2, 2.65, 3.45, 0.5, [[(head, {"bold": True, "color": c})]], size=SUB, margin=0)
    textbox(s, x + 0.2, 3.2, 3.45, 1.0, [body], size=SMALL, margin=0)
code(s, 0.6, 4.55, 12.13, 2.15, '''
[w.upper() for w in words]                 # change every item
[w for w in words if w.startswith("B")]    # keep some items
→ ['Bennet', 'Bennet']
''', size=CODE_S)

# 10 ─ for loops ─────────────────────────────────────────────────────────
s = new_slide("for loops: do something for each item",
              "Read it as: for each of the first three words, call it w, then print w and its length. "
              "The colon and the indentation are compulsory: Python uses indentation, not brackets, to know what "
              "is inside the loop. Colab indents for you after a colon. "
              "A comprehension makes a value; a loop performs an action.")
code(s, 0.6, 1.5, 6.4, 3.3, '''
for w in words[0:3]:
    print(w, len(w))
→ Mr. 3
→ Bennet 6
→ looked 6
''')
textbox(s, 7.4, 1.5, 5.33, 5.2, [
    [("The colon", {"bold": True, "color": DARK, "size": SUB})],
    "ends the first line",
    " ",
    [("The indent", {"bold": True, "color": DARK, "size": SUB})],
    "four spaces: what is repeated",
    " ",
    [("Loop or comprehension?", {"bold": True, "color": DARK, "size": SUB})],
    "Do something: loop. Make a list: comprehension.",
], margin=0, space_after=2)

# 11 ─ Conditions ────────────────────────────────────────────────────────
s = new_slide("Asking questions: True or False",
              "Comparisons give True or False. They appear in the 'if' part of a comprehension and in if statements. "
              "'in' is very useful for text: is this word in this list, is this string inside that one?")
ops = [("==", "equal"), ("!=", "not equal"), ("<  >", "less, greater"), ("<= >=", "or equal to"),
       ("in", "inside?"), ("not in", "absent?")]
for i, (op, d) in enumerate(ops):
    y = 1.5 + i * 0.85
    textbox(s, 0.6, y, 2.0, 0.7, [op], size=28, font=MONO, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    textbox(s, 2.7, y, 3.3, 0.7, [d], anchor=MSO_ANCHOR.MIDDLE, margin=0)
code(s, 6.3, 1.5, 6.43, 5.2, '''
"Bennet" in words
→ True

if len(w) > 5:
    print("long:", w)
else:
    print("short:", w)
''')

# 12 ─ Defining functions ────────────────────────────────────────────────
s = new_slide("Writing your own functions",
              "def defines, return gives back. The name in brackets (w) is a placeholder for whatever you pass in. "
              "Why bother? A function puts a decision, such as what counts as a word, in one named place, "
              "where you can read it, test it and report it in your methods.")
code(s, 0.6, 1.5, 12.13, 2.75, '''
def clean(w):
    return w.strip('.,;:!?"').lower()

clean('"Bennet,')     → 'bennet'
''')
for i, (k, d) in enumerate([("def", "starts a definition"),
                            ("(w)", "the input: whatever is passed in"),
                            (":", "then an indented body"),
                            ("return", "what comes back")]):
    x, y = 0.6 + (i % 2) * 6.15, 4.55 + (i // 2) * 1.1
    textbox(s, x, y, 1.6, 0.8, [k], size=28, font=MONO, color=ACCENT, margin=0, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, x + 1.7, y, 4.3, 0.8, [d], margin=0, anchor=MSO_ANCHOR.MIDDLE)

# 13 ─ Modules and packages ──────────────────────────────────────────────
s = new_slide("Modules, packages, libraries",
              "Vocabulary students will meet in documentation and AI answers. A module is one file of Python code you can "
              "import (re, statistics). A package is a collection of modules, usually installed first (nltk, pandas). "
              "'Library' is informal for either. In Colab most packages are already installed; if not, "
              "!pip install name in a cell. In R the equivalent word is package, loaded with library().")
for i, (term, d, eg) in enumerate([
    ("module", "one file of code", "re"),
    ("package", "a collection of modules", "nltk"),
    ("library", "either, informally", "pandas"),
]):
    y = 1.5 + i * 1.75
    panel(s, 0.6, y, 4.9, 1.5)
    textbox(s, 0.85, y + 0.12, 4.4, 0.5, [[(term, {"bold": True, "color": ACCENT})]], size=SUB, margin=0)
    textbox(s, 0.85, y + 0.7, 4.4, 0.7, [[(d + ": ", {}), (eg, {"font": MONO, "color": DARK})]],
            size=SMALL, margin=0)
code(s, 5.9, 1.5, 6.83, 5.0, '''
import statistics
statistics.mean([3, 6, 6, 2, 4])
→ 4.2

import nltk
nltk.download('punkt_tab')

# in Colab, if a package is missing:
!pip install something
''', size=CODE_S)

# 14 ─ Errors ────────────────────────────────────────────────────────────
s = new_slide("Error messages: read the last line",
              "Errors are normal and informative. Read the last line first; the arrow in the traceback shows which line. "
              "NameError: a typo such as Print or sampel, or a missing import. SyntaxError: a missing quote, bracket "
              "or colon. The most dangerous code is code with no error that gives the wrong answer.")
rows = [("NameError", "unknown name", "Print(x)"),
        ("AttributeError", "no such method", "sample.Upper()"),
        ("IndexError", "past the end of a list", "words[100]"),
        ("TypeError", "wrong type", '"Page " + 5'),
        ("SyntaxError", "not valid Python", 'print("hi)')]
tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(0.6), Inches(1.45), Inches(12.13), Inches(4.2)).table
for c, w in enumerate((3.4, 4.3, 4.43)):
    tbl.columns[c].width = Inches(w)
for r, row in enumerate([("Error", "Means", "Example")] + rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK if r == 0 else (TINT if r % 2 else WHITE)
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        run.font.size = Pt(SMALL)
        run.font.name = MONO if r and c != 1 else BODY
        run.font.bold = r == 0
        run.font.color.rgb = WHITE if r == 0 else (ACCENT if c == 0 else INK)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
textbox(s, 0.6, 5.95, 12.1, 0.8, [[("The worst error has no message: ", {"bold": True, "color": DARK}),
                                   ("wrong answers, quietly.", {"italic": True})]], margin=0)

# 15 ─ Critical habits with AI ───────────────────────────────────────────
s = new_slide("Habits for working with AI-written code",
              "These are the habits we'll practise in the next exercise. Report your decisions: they belong in your "
              "methods, not hidden in code. Lead into the word-frequency prompt.")
habits = [("Test it on something you can check by hand", "a few sentences where you know the answer"),
          ("Ask what it assumed", "what is a word, a sentence, the text?"),
          ("Read its explanation against the code", "explanations can be confidently wrong"),
          ("Ask twice, ask another AI", "same prompt, same answer?"),
          ("Report your decisions", "in your methods, not hidden in code")]
for i, (h, d) in enumerate(habits):
    y = 1.45 + i * 1.06
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(y + 0.1), Inches(0.72), Inches(0.72))
    circ.fill.solid()
    circ.fill.fore_color.rgb = ACCENT
    circ.line.fill.background()
    circ.shadow.inherit = False
    textbox(s, 0.6, y + 0.1, 0.72, 0.72, [[(str(i + 1), {"bold": True, "color": WHITE})]], size=TEXT,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    textbox(s, 1.6, y, 11.1, 0.5, [[(h, {"bold": True, "color": DARK})]], size=SUB, margin=0)
    textbox(s, 1.6, y + 0.52, 11.1, 0.45, [d], size=SMALL, color=MUTED, margin=0)

# 16 ─ Exercise ──────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
textbox(s, 0.9, 0.6, W - 1.8, 1.1, ["Your turn: ask an AI"], size=HUGE, font=HEAD, color=DARK,
        anchor=MSO_ANCHOR.MIDDLE, margin=0)
textbox(s, 0.9, 1.95, W - 1.8, 2.4, [[("“Write Python code for Google Colab that downloads Pride and Prejudice "
                                       "from Project Gutenberg and lists the 20 most common words.”",
                                       {"italic": True, "color": DARK})]], size=32, margin=0)
textbox(s, 0.9, 4.4, W - 1.8, 2.3, ["Does it run? What did it read? What counts as a word? "
                                    "Did it remove anything? How many times does it count “pride”?"],
        size=28, margin=0)
footer(s)
notes(s, "Deliberately vague prompt. Compare with the concordance count for 'pride' (56 in the full download). "
         "Look for: Gutenberg licence left in; case; punctuation and curly quotes; underscores for italics; "
         "stop words silently removed (the very words stylometry needs). Compare top-20 lists across AIs in the room.")

out = sys.argv[1] if len(sys.argv) > 1 else "slides/python_syntax_slides.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides")
